#!/usr/bin/env python3
"""
Offline corruption/regression checker for xlsx files XLSV has written.

Uses openpyxl -- a mature, independently-developed xlsx library with no code
in common with this app's own hand-rolled XML splicing (service.swift) -- as
an outside opinion. The app's own XMLValidator only confirms the XML is
well-formed; it can't catch a file that's well-formed but semantically wrong
(duplicate row numbers, a dimension that doesn't match the data, a mergeCell
range pointing past the sheet, a shared-string index out of bounds, etc.).
That's exactly the class of bug a from-scratch XML splice (row/col delete,
the <cols>/<row ht=> patching added 2026-08-23) is most likely to introduce.

Two modes:

  Single-file sanity check -- opens the file and runs structural checks:
      python3 xlsx_corruption_check.py path/to/file.xlsx

  Before/after regression check -- also diffs cell values between an
  original file and a saved-after-edit copy, flagging any cell that changed
  content unexpectedly (use --expect-changed to allow specific cells, e.g.
  the ones you actually meant to edit):
      python3 xlsx_corruption_check.py before.xlsx after.xlsx
      python3 xlsx_corruption_check.py before.xlsx after.xlsx --expect-changed A1,B2

Exit code is non-zero if any check fails, so this is CI/pre-commit friendly.
"""

import argparse
import re
import sys
import zipfile

try:
    import openpyxl
    from openpyxl.utils import range_boundaries
except ImportError:
    print("openpyxl is required: pip3 install openpyxl", file=sys.stderr)
    sys.exit(2)


def load(path):
    """Opens with openpyxl; a raised exception here IS the corruption signal --
    real Excel/Sheets would refuse this file too."""
    try:
        return openpyxl.load_workbook(path, data_only=False), []
    except Exception as e:
        return None, [f"FAILED TO OPEN: {type(e).__name__}: {e}"]


def raw_sheet_xml(path):
    """Reads each xl/worksheets/sheetN.xml's raw text directly from the zip,
    bypassing openpyxl entirely. Needed because openpyxl normalizes/tolerates
    exactly the kinds of things a botched row/col delete would produce --
    confirmed both ws.dimensions/max_row/max_column (silently inflated to
    include any merged-cell range) and ws.iter_rows() (silently coalesces two
    <row r="N"> elements sharing the same N into one logical row) hide rather
    than surface the underlying malformed markup.

    Returns {sheet_index (1-based, from the "sheetN.xml" filename): raw XML text}.
    Pairing this index back to a specific worksheet name is done by naive
    ascending order against wb.worksheets in check_structure -- good enough
    for sanity-checking this app's own output (sheetN.xml numbering already
    matches its own creation-order convention), not a general-purpose
    workbook.xml.rels-based mapping.
    """
    sheets = {}
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            m = re.match(r"xl/worksheets/sheet(\d+)\.xml$", name)
            if not m:
                continue
            sheets[int(m.group(1))] = z.read(name).decode("utf-8", errors="replace")
    return sheets


def parse_dimension(xml):
    dim_match = re.search(r'<dimension ref="([^"]+)"', xml)
    if not dim_match:
        return None
    ref = dim_match.group(1)
    try:
        min_col, min_row, max_col, max_row = range_boundaries(ref if ":" in ref else f"{ref}:{ref}")
        return (max_row, max_col)
    except ValueError:
        return None


def duplicate_row_numbers(xml):
    row_numbers = [int(n) for n in re.findall(r'<row\b[^>]*\br="(\d+)"', xml)]
    seen, dupes = set(), set()
    for n in row_numbers:
        if n in seen:
            dupes.add(n)
        seen.add(n)
    return sorted(dupes)


def check_structure(wb, path):
    """Sanity checks openpyxl's own decode doesn't already guarantee --
    things that could still be internally inconsistent even in a file that
    opens without raising."""
    problems = []
    if not wb.sheetnames:
        problems.append("workbook has zero sheets")
        return problems

    raw_sheets = raw_sheet_xml(path)
    raw_sheets_in_order = [raw_sheets[i] for i in sorted(raw_sheets.keys())]

    for sheet_pos, ws in enumerate(wb.worksheets):
        raw_xml = raw_sheets_in_order[sheet_pos] if sheet_pos < len(raw_sheets_in_order) else None
        raw_dim = parse_dimension(raw_xml) if raw_xml is not None else None
        if raw_xml is None:
            problems.append(f"sheet '{ws.title}': couldn't locate this sheet's own worksheet XML in the zip")
        elif raw_dim is None:
            problems.append(f"sheet '{ws.title}': missing/unparseable <dimension ref=> in its worksheet XML")

        # Row numbers must be unique in the raw XML -- a delete that failed
        # to renumber correctly, or a splice that duplicated a <row>, shows
        # up here as a repeat. Checked against the raw markup, not
        # ws.iter_rows(), since openpyxl silently coalesces two <row r="N">
        # elements sharing the same N into a single logical row rather than
        # surfacing the duplicate.
        if raw_xml is not None:
            for r in duplicate_row_numbers(raw_xml):
                problems.append(f"sheet '{ws.title}': row {r} appears more than once in the raw XML")

        # mergeCell ranges must stay within the sheet's own RAW declared
        # dimension -- exactly the kind of thing a row/col delete that
        # doesn't adjust merges (or update <dimension>) would break. Checked
        # against the raw XML value, not openpyxl's own max_row/max_column,
        # since the latter already silently expands to cover any merge.
        if raw_dim is not None:
            raw_max_row, raw_max_col = raw_dim
            for merged_range in ws.merged_cells.ranges:
                if merged_range.max_row > raw_max_row or merged_range.max_col > raw_max_col:
                    problems.append(
                        f"sheet '{ws.title}': merged range {merged_range} extends past the sheet's "
                        f"own declared <dimension> (max_row={raw_max_row}, max_col={raw_max_col})"
                    )

        # Column widths / row heights set via customWidth/customHeight --
        # not a failure on their own, just reported so a size-patch round
        # trip is easy to eyeball.
        sized_cols = {letter: dim.width for letter, dim in ws.column_dimensions.items() if dim.width is not None}
        sized_rows = {idx: dim.height for idx, dim in ws.row_dimensions.items() if dim.height is not None}
        if sized_cols:
            print(f"  sheet '{ws.title}': custom column widths -> {sized_cols}")
        if sized_rows:
            print(f"  sheet '{ws.title}': custom row heights -> {sized_rows}")

    return problems


def diff_cells(before_path, after_path, expect_changed):
    """Compares every cell's value between two workbooks sheet-by-sheet,
    flagging any change not explicitly allow-listed via --expect-changed.
    Doesn't compare styles/formatting -- that's a much fuzzier equality
    check and out of scope for a corruption check specifically."""
    problems = []
    wb_before = openpyxl.load_workbook(before_path, data_only=False)
    wb_after = openpyxl.load_workbook(after_path, data_only=False)

    common_sheets = set(wb_before.sheetnames) & set(wb_after.sheetnames)
    for name in wb_before.sheetnames:
        if name not in wb_after.sheetnames:
            problems.append(f"sheet '{name}' present before, missing after")
    for name in wb_after.sheetnames:
        if name not in wb_before.sheetnames:
            problems.append(f"sheet '{name}' present after, missing before (unexpected new sheet)")

    for name in common_sheets:
        ws_before = wb_before[name]
        ws_after = wb_after[name]
        max_row = max(ws_before.max_row, ws_after.max_row)
        max_col = max(ws_before.max_column, ws_after.max_column)
        for row in range(1, max_row + 1):
            for col in range(1, max_col + 1):
                cell_before = ws_before.cell(row=row, column=col)
                cell_after = ws_after.cell(row=row, column=col)
                if cell_before.value != cell_after.value:
                    ref = cell_before.coordinate
                    if ref in expect_changed:
                        continue
                    problems.append(
                        f"sheet '{name}' {ref}: value changed unexpectedly "
                        f"({cell_before.value!r} -> {cell_after.value!r})"
                    )
    return problems


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("file", help="xlsx file to check (or the 'before' file in before/after mode)")
    parser.add_argument("after_file", nargs="?", help="the 'after' file, for before/after regression mode")
    parser.add_argument("--expect-changed", default="", help="comma-separated cell refs allowed to differ, e.g. A1,B2")
    args = parser.parse_args()

    all_problems = []

    print(f"Opening {args.file} ...")
    wb, open_problems = load(args.file)
    all_problems += open_problems
    if wb:
        all_problems += check_structure(wb, args.file)

    if args.after_file:
        print(f"Opening {args.after_file} ...")
        wb_after, open_problems_after = load(args.after_file)
        all_problems += open_problems_after
        if wb_after:
            all_problems += check_structure(wb_after, args.after_file)

        if wb and wb_after:
            expect_changed = {c.strip().upper() for c in args.expect_changed.split(",") if c.strip()}
            print(f"Diffing cell values ({args.file} -> {args.after_file}) ...")
            all_problems += diff_cells(args.file, args.after_file, expect_changed)

    print()
    if all_problems:
        print(f"FAIL -- {len(all_problems)} problem(s):")
        for p in all_problems:
            print(f"  - {p}")
        sys.exit(1)
    else:
        print("PASS -- no corruption/regression detected.")
        sys.exit(0)


if __name__ == "__main__":
    main()
